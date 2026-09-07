"""뉴스 목록을 보기 쉬운 Excel 파일로 저장합니다."""

from __future__ import annotations

from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = ("제목", "등록일", "링크")


def write_news_to_excel(news_items: Sequence[dict[str, str]], output_path: Path) -> None:
    """뉴스 데이터를 ``output_path``에 xlsx 파일로 저장합니다."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "USMEF 뉴스"
    worksheet.append(HEADERS)

    for item in news_items:
        worksheet.append([item["제목"], item["등록일"], item["링크"]])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 60
    worksheet.column_dimensions["B"].width = 14
    worksheet.column_dimensions["C"].width = 95

    for row in worksheet.iter_rows(min_row=2):
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        row[1].alignment = Alignment(horizontal="center", vertical="top")
        row[2].alignment = Alignment(vertical="top")
        row[2].hyperlink = row[2].value
        row[2].style = "Hyperlink"

    worksheet.row_dimensions[1].height = 22
    for column_index in range(1, len(HEADERS) + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].bestFit = False

    workbook.save(output_path)
